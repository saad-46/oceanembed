"""Build the OceanSight SIH 2026 idea-submission deck by editing the official template.

    python scripts/build_sih_deck.py --template <SIH2026-IDEA-Presentation-Format.pptx>
        --shots deliverables/assets --out deliverables/OceanSight_SIH2026_Idea_Submission.pptx

``--template`` may also be a previously built 6-slide deck: its template parts (SIH logo, title
placeholders, team oval, blue footer, page numbers) are kept and every builder-added shape is
rebuilt, so the design can be iterated without the original template file.

Screenshots come from the running prototype (headless Chrome, 1600x900 @2x):
/map?date=2023-05-11&depth=100 and /analysis?mode=cyclone. Icons are Lucide line icons
(ISC licence) pre-rendered to deliverables/assets/icons; pass ``--lucide <lucide-static/icons>``
to render any that are missing.

Keeps the template's slides, SIH logo, title placeholders, team oval, blue footer and page
numbers; replaces only the content text boxes; deletes the "IMPORTANT INSTRUCTIONS" slide.
Every number comes from docs/RESULTS.md (computed by the prototype pipeline).
"""
from __future__ import annotations

import argparse
import re
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

NAVY = RGBColor(0x0B, 0x2F, 0x4E)
BLUE = RGBColor(0x00, 0x70, 0xC0)      # SIH footer blue
TEAL = RGBColor(0x0E, 0x86, 0x9A)
TEAL_LIGHT = RGBColor(0xE3, 0xF3, 0xF6)
TEAL_LINE = RGBColor(0x9F, 0xD0, 0xD8)
AQUA = RGBColor(0x7F, 0xD3, 0xE0)      # accent on navy
PANEL = RGBColor(0xF2, 0xF6, 0xF9)
SKY = RGBColor(0xED, 0xF4, 0xFA)       # section panel tint
SKY_LINE = RGBColor(0xA9, 0xC6, 0xE3)
LINE = RGBColor(0xC9, 0xD6, 0xE2)
TEXT = RGBColor(0x1B, 0x23, 0x30)
MUTED = RGBColor(0x55, 0x65, 0x78)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x1B, 0x8A, 0x4B)
GREEN_LIGHT = RGBColor(0xEA, 0xF5, 0xEE)
GREEN_LINE = RGBColor(0xA9, 0xD3, 0xB7)
AMBER = RGBColor(0xB3, 0x6B, 0x00)
AMBER_LIGHT = RGBColor(0xFD, 0xF3, 0xE6)
AMBER_LINE = RGBColor(0xE9, 0xC0, 0x8A)
ICE = RGBColor(0xBF, 0xE3, 0xEA)
FONT = "Arial"
TEAM_NAME = "CodeCrafters"
TEAM_ID = "135494"

ICON_DIR: Path | None = None       # set in main(): <shots>/icons
LUCIDE_DIR: Path | None = None     # optional source SVGs
TMP = Path(tempfile.mkdtemp(prefix="sih_deck_"))


# ---------------------------------------------------------------- helpers
def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None, line_w=0.75):
    """radius is the corner radius in inches (rounded rectangles only)."""
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    unstyle(s)
    if s.has_text_frame:
        s.text_frame.text = ""
    return s


def unstyle(shape):
    """No shadow, and drop the theme style reference (its effectRef adds a shadow in some renderers)."""
    shape.shadow.inherit = False
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def card(slide, x, y, w, h, fill=WHITE, line=LINE, radius=0.08, line_w=1.0):
    return box(slide, x, y, w, h, fill=fill, line=line, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius, line_w=line_w)


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, margin=0.04):
    """paras: list of paragraphs; each paragraph = dict(runs=[(text, {fmt})], align, space_after, bullet)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    fill_paras(tf, paras)
    return tb


def set_font(run, name=FONT):
    """Set latin + east-asian + complex-script typefaces so symbols like ° and σ don't fall back."""
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def fill_paras(tf, paras):
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", PP_ALIGN.LEFT)
        if "space_after" in p:
            para.space_after = Pt(p["space_after"])
        if "space_before" in p:
            para.space_before = Pt(p["space_before"])
        para.line_spacing = p.get("line_spacing", 1.0)
        for t, f in p["runs"]:
            r = para.add_run()
            r.text = t
            set_font(r, f.get("font", FONT))
            r.font.size = Pt(f.get("size", 12))
            r.font.bold = f.get("bold", False)
            r.font.italic = f.get("italic", False)
            r.font.color.rgb = f.get("color", TEXT)
            if f.get("link"):
                r.hyperlink.address = f["link"]
                r.font.color.rgb = f.get("color", TEXT)


def P(*runs, align=PP_ALIGN.LEFT, **kw):
    return {"runs": list(runs), "align": align, **kw}


def lines(t, size, color=TEXT, bold=False, align=PP_ALIGN.CENTER):
    """One paragraph per line, so every line of a centred caption is centred."""
    return [P(R(ln, size, bold, color), align=align) for ln in t.split("\n")]


def R(t, size=12, bold=False, color=TEXT, italic=False, font=FONT, link=None):
    return (t, {"size": size, "bold": bold, "color": color, "italic": italic, "font": font, "link": link})


def icon_png(name: str, color: RGBColor) -> Path:
    """Lucide line icon rendered as a transparent PNG in the given colour (cached in ICON_DIR)."""
    hexc = str(color)
    out = ICON_DIR / f"{name}-{hexc}.png"
    if out.exists():
        return out
    if LUCIDE_DIR is None:
        raise SystemExit(f"icon {out.name} missing; pass --lucide <lucide-static/icons> to render it")
    import cairosvg  # only needed when rendering new icons

    svg = (LUCIDE_DIR / f"{name}.svg").read_text()
    svg = svg.replace("currentColor", f"#{hexc}")
    svg = re.sub(r'stroke-width="[\d.]+"', 'stroke-width="2"', svg)
    ICON_DIR.mkdir(parents=True, exist_ok=True)
    cairosvg.svg2png(bytestring=svg.encode(), write_to=str(out), output_width=256, output_height=256)
    return out


def icon(slide, name, x, y, size, color=TEAL):
    return slide.shapes.add_picture(str(icon_png(name, color)), Inches(x), Inches(y), Inches(size), Inches(size))


def badge(slide, name, x, y, size, bg=TEAL, fg=WHITE, round_=False, line=None):
    """Icon centred in a filled rounded square (or circle) — the deck's recurring motif."""
    if round_:
        box(slide, x, y, size, size, fill=bg, line=line, shape=MSO_SHAPE.OVAL)
    else:
        box(slide, x, y, size, size, fill=bg, line=line, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=size * 0.22)
    pad = size * 0.2
    icon(slide, name, x + pad, y + pad, size - 2 * pad, fg)


def section_label(slide, x, y, w, label, sub=None, ic="circle-check", color=NAVY, bg=TEAL, size=13):
    """Section heading: icon badge + bold heading (+ muted sub-heading)."""
    badge(slide, ic, x, y, 0.3, bg=bg)
    runs = [R(label, size, True, color)]
    if sub:
        runs.append(R("   " + sub, 10, False, MUTED))
    text(slide, x + 0.38, y - 0.02, w - 0.38, 0.34, [P(*runs)], anchor=MSO_ANCHOR.MIDDLE, margin=0.01)


def arrow(slide, x, y, w=0.16, h=0.22, color=TEAL):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid()
    a.fill.fore_color.rgb = color
    a.line.fill.background()
    unstyle(a)
    return a


def chevron(slide, x, y, w=0.11, h=0.2, color=TEAL):
    a = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid()
    a.fill.fore_color.rgb = color
    a.line.fill.background()
    unstyle(a)
    return a


def picture(slide, path, x, y, w=None, h=None, crop=None):
    """crop = (left, top, right, bottom) fractions of the source image."""
    img = Image.open(path)
    if crop:
        W, H = img.size
        img = img.crop((int(crop[0] * W), int(crop[1] * H), int(crop[2] * W), int(crop[3] * H)))
        path = TMP / (Path(path).stem + f"_crop_{abs(hash(crop)) % 10**6}.png")
        img.save(path)
    ar = img.size[0] / img.size[1]
    if w is not None and h is None:
        h = w / ar
    elif h is not None and w is None:
        w = h * ar
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    pic.line.color.rgb = LINE
    pic.line.width = Pt(0.75)
    return pic, w, h


def text_w(t, size, bold=False):
    """Rough rendered width (inches) of Arial text — for centring an icon next to a label."""
    return len(t) * size / 72 * (0.56 if bold else 0.5)


def remove_shape(shape):
    el = shape._element
    el.getparent().remove(el)


def delete_slide(prs, index):
    sldIdLst = prs.slides._sldIdLst
    sld = sldIdLst[index]
    prs.part.drop_rel(sld.rId)
    sldIdLst.remove(sld)


def set_title(slide, lines):
    """Replace the title placeholder text, keeping its font family (Times New Roman) and position."""
    ph = slide.shapes.title
    tf = ph.text_frame
    first = tf.paragraphs[0]
    font_name = next((r.font.name for r in first.runs if r.font.name), "Times New Roman")
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(first.runs):
        r._r.getparent().remove(r._r)
    for i, (t, size, color, bold) in enumerate(lines):
        para = first if i == 0 else tf.add_paragraph()
        para.alignment = PP_ALIGN.CENTER
        r = para.add_run()
        r.text = t
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        set_font(r, font_name if i == 0 else FONT)


def set_team_oval(slide):
    """Put the team name into the template's "Your Team Name" oval (keeps its outline styling)."""
    for sh in slide.shapes:
        if sh.has_text_frame and "Team" in sh.text_frame.text and sh.name.startswith("Oval"):
            tf = sh.text_frame
            tf.clear()
            tf.word_wrap = False
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            for side in ("left", "right", "top", "bottom"):
                setattr(tf, f"margin_{side}", Inches(0.02))
            para = tf.paragraphs[0]
            para.alignment = PP_ALIGN.CENTER
            r = para.add_run()
            r.text = TEAM_NAME
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = NAVY
            set_font(r)


def is_template_part(sh) -> bool:
    """Template furniture on content slides: placeholders, team oval, SIH logo, blue footer bar."""
    if sh.is_placeholder or sh.name.startswith("Oval"):
        return True
    x, y, w = sh.left / 914400, sh.top / 914400, sh.width / 914400
    if sh.shape_type == 13 and x > 10.0 and y < 0.2:      # SIH logo, top right
        return True
    return sh.shape_type == 1 and y > 6.8 and w > 13.0    # footer bar


def clear_content_boxes(slide, rebuilt: bool):
    for sh in list(slide.shapes):
        if rebuilt:
            if not is_template_part(sh):
                remove_shape(sh)
        elif sh.shape_type == 17 and sh.has_text_frame:  # TEXT_BOX with template pointer text
            remove_shape(sh)


# ---------------------------------------------------------------- slides
def slide1(s):
    # Subtitle placeholder ("TITLE PAGE") -> project name, placed left so it doesn't cover the SIH graphic
    sub = next(sh for sh in s.shapes if sh.is_placeholder and sh.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE)
    sub.left, sub.top, sub.width, sub.height = Inches(0.36), Inches(1.12), Inches(7.1), Inches(1.25)
    tf = sub.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.TOP
    fill_paras(tf, [
        P(R("OceanSight", 48, True, NAVY, font="Arial"), space_after=2),
        P(R("Seeing 1000 m beneath the ocean surface — from satellites alone", 15, True, TEAL)),
    ])
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.LEFT
    # Fields text box (keeps its position and bullets style)
    tb = next(sh for sh in s.shapes if sh.shape_type == 17 and "Problem Statement ID" in sh.text_frame.text)
    tb.left, tb.top, tb.width, tb.height = Inches(0.36), Inches(2.62), Inches(6.6), Inches(4.6)
    tf = tb.text_frame
    tf.word_wrap = True
    paras = [p for p in tf.paragraphs if p.text.strip()]
    values = [
        ("Problem Statement ID – ", "SIH26066"),
        ("Problem Statement Title – ", "OceanEmbed – Satellite Embedding-Based Deep Learning Framework for "
                                       "Reconstruction of Subsurface Ocean Temperature from Surface Satellite Observations"),
        ("Theme – ", "Space Technology"),
        ("PS Category – ", "Software"),
        ("Team ID – ", TEAM_ID),
        ("Team Name (Registered on portal) – ", TEAM_NAME),
    ]
    for para, (label, value) in zip(paras, values):
        for r in list(para.runs):
            r._r.getparent().remove(r._r)
        para.space_after = Pt(14)
        para.space_before = Pt(0)
        para.line_spacing = 1.05
        para.alignment = PP_ALIGN.LEFT
        r1 = para.add_run(); r1.text = label
        r1.font.size = Pt(17); r1.font.bold = True; r1.font.color.rgb = TEXT; set_font(r1)
        r2 = para.add_run(); r2.text = value
        small = len(value) > 40
        r2.font.size = Pt(15 if small else 17); r2.font.bold = not small
        r2.font.color.rgb = NAVY; set_font(r2)
    # drop empty paragraphs (template has blank lines)
    for p in list(tf.paragraphs):
        if not p.text.strip():
            p._p.getparent().remove(p._p)


def slide2(s, shots, rebuilt):
    clear_content_boxes(s, rebuilt)
    set_title(s, [("OceanSight", 32, NAVY, True),
                  ("Daily 0–1000 m ocean temperature for the North Indian Ocean — from satellites alone", 13, TEAL, False)])
    L, W = 0.35, 8.5
    # A. Proposed solution
    ay, ah = 1.32, 1.06
    card(s, L, ay, W, ah, fill=SKY, line=SKY_LINE, radius=0.1)
    section_label(s, L + 0.14, ay + 0.09, W - 0.28, "Proposed Solution", "Detailed explanation of the proposed solution",
                  ic="lightbulb")
    text(s, L + 0.14, ay + 0.42, W - 0.28, 0.6, [P(
        R("OceanSight ", 11, True, NAVY),
        R("reconstructs ocean temperature at ", 11),
        R("15 depths (0–1000 m), every day, on a 0.25° grid", 11, True),
        R(" over the Bay of Bengal & Arabian Sea — using only ", 11),
        R("5 satellite surface fields", 11, True),
        R(" — and turns it into cyclone-heat and mixed-layer products for ocean analysts and forecasters.", 11),
        line_spacing=1.0)], margin=0.02)
    # B. How it addresses the problem
    by, bh = 2.48, 1.94
    card(s, L, by, W, bh, fill=SKY, line=SKY_LINE, radius=0.1)
    section_label(s, L + 0.14, by + 0.09, W - 0.28, "How it addresses the problem", ic="target")
    gy, gh = by + 0.44, 0.54
    gx, gw = L + 0.14, 3.0
    card(s, gx, gy, gw, gh, fill=AMBER_LIGHT, line=AMBER_LINE, radius=0.08)
    icon(s, "triangle-alert", gx + 0.1, gy + (gh - 0.26) / 2, 0.26, AMBER)
    text(s, gx + 0.42, gy, gw - 0.48, gh, [
        P(R("THE GAP", 9, True, AMBER)),
        P(R("Floats, buoys & ships sample the subsurface only at scattered points and times", 9))],
        anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
    arrow(s, gx + gw + 0.08, gy + gh / 2 - 0.11, 0.26, 0.22)
    ox = gx + gw + 0.42
    ow = L + W - 0.14 - ox
    card(s, ox, gy, ow, gh, fill=WHITE, line=TEAL, radius=0.08, line_w=1.25)
    icon(s, "satellite", ox + 0.1, gy + (gh - 0.26) / 2, 0.26, TEAL)
    text(s, ox + 0.42, gy, ow - 0.48, gh, [
        P(R("OCEANSIGHT", 9, True, TEAL)),
        P(R("Satellites see the whole basin daily; a deep network learns the physical surface → subsurface link "
            "(e.g. sea-level anomaly → thermocline depth)", 9))], anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
    stages = [
        ("satellite", "Satellite inputs", "SST · SSS · SLA ·\ncurrents · winds"),
        ("grid-3x3", "Harmonise", "QC + regrid to\n0.25°, daily"),
        ("brain-circuit", "Satellite embedding", "U-Net encoder\ncompresses the basin"),
        ("layers", "Reconstruct", "15 depths, 0–1000 m\n+ uncertainty"),
        ("gauge", "Derive", "TCHP · MLD ·\nD20 · D26"),
        ("monitor", "Deliver", "GIS dashboard ·\nAPI · reports"),
    ]
    n, gap = len(stages), 0.16
    iw = W - 0.28
    bw = (iw - gap * (n - 1)) / n
    y0, sh = gy + gh + 0.1, 0.76
    for i, (ic, t, sub) in enumerate(stages):
        x = L + 0.14 + i * (bw + gap)
        core = i == 2
        card(s, x, y0, bw, sh, fill=TEAL if core else WHITE, line=TEAL if core else LINE, radius=0.08)
        icon(s, ic, x + (bw - 0.22) / 2, y0 + 0.05, 0.22, WHITE if core else TEAL)
        text(s, x, y0 + 0.28, bw, 0.18, [P(R(t, 8.5, True, WHITE if core else NAVY), align=PP_ALIGN.CENTER)],
             anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
        text(s, x, y0 + 0.44, bw, 0.3, lines(sub, 8, WHITE if core else TEXT),
             anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
        if i < n - 1:
            chevron(s, x + bw + (gap - 0.09) / 2, y0 + sh / 2 - 0.09, 0.09, 0.18)
    # Right: prototype screenshot
    px, pw = 9.05, 3.95
    card(s, px, 1.32, pw, by + bh - 1.32, fill=PANEL, line=LINE, radius=0.1)
    _, _, ph = picture(s, shots / "map.png", px + 0.1, 1.42, w=pw - 0.2, crop=(0.0, 0.14, 0.72, 0.845))
    text(s, px + 0.1, 1.42 + ph + 0.1, pw - 0.2, by + bh - 1.42 - ph - 0.2, [P(
        R("Working prototype: ", 8.5, True, NAVY),
        R("reconstructed temperature at 100 m on 11 May 2023 (held-out year) — dark purple ≈ 17 °C to yellow ≈ 28 °C; "
          "green dots = Argo floats; white line = Cyclone Mocha track (IBTrACS)", 8.5, False, MUTED))],
        anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
    # C. Innovation & uniqueness — the slide's focal band
    top, bandh = 4.52, 2.36
    box(s, 0.3, top, 12.73, bandh, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    badge(s, "sparkles", 0.5, top + 0.1, 0.34, bg=TEAL)
    text(s, 0.94, top + 0.08, 11.9, 0.38, [P(
        R("INNOVATION AND UNIQUENESS OF THE SOLUTION", 14, True, WHITE),
        R("   —   why OceanSight is different", 11, False, AQUA))], anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
    cards = [
        ("badge-check", "01", "Validated beyond training data",
         "Scored against real Argo floats from a year it never saw — next to the training product's own error ceiling "
         "and an independent EN4 cross-check.",
         "2,639", "held-out Argo profiles (2023)"),
        ("crosshair", "02", "Calibrated to the real ocean",
         "Every profile carries an error bar calibrated on 2022 floats that holds on 2023 — analysts can see when not to trust a value.",
         "70%", "of Argo values within ±1σ (ideal 68%)"),
        ("tornado", "03", "Reconstruction → decisions",
         "Daily cyclone heat potential, mixed-layer and isotherm depths at 0.25°, replayed on real IBTrACS cyclone tracks.",
         "TCHP · MLD", "D20 · D26 — every day, 2019–2023"),
    ]
    cg = 0.18
    cw = (12.33 - 2 * cg) / 3
    cy, ch = top + 0.52, 1.46
    for i, (ic, num, title, body, metric, mlabel) in enumerate(cards):
        x = 0.5 + i * (cw + cg)
        card(s, x, cy, cw, ch, fill=WHITE, line=None, radius=0.1)
        badge(s, ic, x + 0.14, cy + 0.1, 0.44, bg=TEAL_LIGHT, fg=TEAL, round_=True)
        text(s, x + 0.68, cy + 0.08, cw - 0.8, 0.2, [P(R(f"USP {num}", 9, True, TEAL))], anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
        text(s, x + 0.68, cy + 0.27, cw - 0.8, 0.28, [P(R(title, 12.5, True, NAVY))], anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
        text(s, x + 0.14, cy + 0.58, cw - 0.28, 0.46, [P(R(body, 9, False, TEXT), line_spacing=1.0)], margin=0.01)
        box(s, x + 0.1, cy + ch - 0.38, cw - 0.2, 0.3, fill=TEAL_LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
        text(s, x + 0.2, cy + ch - 0.38, cw - 0.4, 0.3, [P(R(metric + "  ", 15, True, TEAL), R(mlabel, 9, False, MUTED))],
             anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
    text(s, 0.5, top + bandh - 0.34, 12.33, 0.28, [P(
        R("Satellite surface fields  →  validated, uncertainty-aware 3-D ocean temperature  →  INCOIS decision products",
          11, True, ICE, italic=True),
        align=PP_ALIGN.CENTER)], anchor=MSO_ANCHOR.MIDDLE, margin=0.01)


def slide3(s, rebuilt):
    clear_content_boxes(s, rebuilt)
    # Left: technologies
    LX, LW = 0.35, 4.35
    section_label(s, LX, 1.30, LW, "Technologies to be used", ic="cpu")
    rows = [
        ("satellite", "Satellite data", "NOAA OISST (SST) · SMAP + SMOS (SSS) · NOAA altimetry (SLA, currents) · NCEI Seawinds"),
        ("badge-check", "Target & validation", "HYCOM 1/12° analysis (GLORYS adapter built) · Argo via argopy · Met Office EN4"),
        ("brain-circuit", "ML / AI", "PyTorch U-Net (satellite embedding) · LightGBM baseline"),
        ("database", "Data engineering", "Python · xarray · Zarr (daily cubes) · NumPy regridding"),
        ("map", "Geospatial / DB", "PostgreSQL + PostGIS (spatial indexes, nearest-float queries)"),
        ("server", "Backend", "FastAPI — typed REST API, PDF/CSV reports"),
        ("monitor", "Frontend / GIS", "Next.js · MapLibre GL · deck.gl · Recharts"),
        ("wrench", "Engineering", "Docker · GitHub Actions CI · 44 automated tests · CPU-only"),
    ]
    y, rh, pitch = 1.7, 0.57, 0.648
    for ic, cat, tools in rows:
        card(s, LX, y, LW, rh, fill=PANEL, line=LINE, radius=0.08)
        badge(s, ic, LX + 0.1, y + (rh - 0.34) / 2, 0.34, bg=NAVY)
        text(s, LX + 0.52, y, 1.2, rh, [P(R(cat, 9.3, True, NAVY))], anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
        text(s, LX + 1.74, y, LW - 1.82, rh, [P(R(tools, 9, False, TEXT))], anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
        y += pitch
    # Right: methodology
    RX, RW = 4.95, 8.05
    section_label(s, RX, 1.30, RW, "Methodology and process for implementation", ic="workflow")
    stages = [
        ("Satellite data", "5 fields · 1,826 days\n(2019–2023)"),
        ("Harmonise", "QC · gap-fill ·\nregrid to 0.25°"),
        ("Features", "7 channels + lat/lon\n+ season"),
        ("U-Net model", "embedding → 15 depths\n+ uncertainty (σ)"),
        ("Validate", "held-out Argo · EN4\n· σ calibration"),
        ("Serve", "Zarr + PostGIS +\nFastAPI → GIS app"),
    ]
    n, gap = len(stages), 0.14
    bw = (RW - gap * (n - 1)) / n
    y0, sh = 1.7, 1.06
    for i, (t, sub) in enumerate(stages):
        x = RX + i * (bw + gap)
        core = i == 3
        card(s, x, y0, bw, sh, fill=TEAL if core else WHITE, line=TEAL if core else LINE, radius=0.08)
        box(s, x + (bw - 0.28) / 2, y0 + 0.08, 0.28, 0.28, fill=WHITE if core else NAVY, shape=MSO_SHAPE.OVAL)
        text(s, x + (bw - 0.28) / 2, y0 + 0.08, 0.28, 0.28, [P(R(str(i + 1), 10, True, TEAL if core else WHITE),
                                                            align=PP_ALIGN.CENTER)], anchor=MSO_ANCHOR.MIDDLE, margin=0)
        text(s, x, y0 + 0.4, bw, 0.2, [P(R(t, 9.5, True, WHITE if core else NAVY), align=PP_ALIGN.CENTER)],
             anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
        text(s, x, y0 + 0.6, bw, 0.42, lines(sub, 8, WHITE if core else TEXT),
             anchor=MSO_ANCHOR.MIDDLE, margin=0)
        if i < n - 1:
            chevron(s, x + bw + (gap - 0.08) / 2, y0 + sh / 2 - 0.09, 0.08, 0.18)
    # Year split timeline
    ty = 2.9
    icon(s, "calendar-range", RX, ty + 0.03, 0.22, TEAL)
    text(s, RX + 0.3, ty, RW - 0.3, 0.28, [P(R("Honest evaluation — split by whole years, never random days", 10.5, True, NAVY))],
         anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
    segs = [("2019", NAVY), ("2020", NAVY), ("2021", NAVY), ("2022", TEAL), ("2023", AMBER)]
    sw = (RW - 0.1 * 4) / 5
    for i, (yr, col) in enumerate(segs):
        x = RX + i * (sw + 0.1)
        box(s, x, ty + 0.36, sw, 0.36, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
        text(s, x, ty + 0.36, sw, 0.36, [P(R(yr, 10.5, True, WHITE), align=PP_ALIGN.CENTER)], anchor=MSO_ANCHOR.MIDDLE)
    text(s, RX, ty + 0.75, 3 * sw + 0.2, 0.26, [P(R("TRAIN — model fitting, normalisation, climatology", 9, True, NAVY), align=PP_ALIGN.CENTER)])
    text(s, RX + 3 * (sw + 0.1), ty + 0.75, sw, 0.26, [P(R("VALIDATE — early stop", 9, True, TEAL), align=PP_ALIGN.CENTER)])
    text(s, RX + 4 * (sw + 0.1), ty + 0.75, sw, 0.26, [P(R("TEST — fully held out", 9, True, AMBER), align=PP_ALIGN.CENTER)])
    # Model ladder
    my = 4.02
    icon(s, "scale", RX, my + 0.03, 0.22, TEAL)
    text(s, RX + 0.3, my, RW - 0.3, 0.28, [P(R("Models, compared side by side — the deep model has to earn its place", 10.5, True, NAVY))],
         anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
    models = [
        ("chart-line", "Seasonal climatology", "harmonic fit per cell & depth — the floor to beat"),
        ("network", "LightGBM", "one model per depth — fast, interpretable baseline"),
        ("brain-circuit", "U-Net encoder–decoder", "bottleneck = satellite embedding; 1.09 M parameters; mean + σ per depth"),
    ]
    mg = 0.3
    mw = (RW - 2 * mg) / 3
    for i, (ic, t, sub) in enumerate(models):
        x = RX + i * (mw + mg)
        best = i == 2
        card(s, x, my + 0.36, mw, 0.86, fill=TEAL_LIGHT if best else WHITE, line=TEAL if best else LINE,
             radius=0.08, line_w=1.5 if best else 1.0)
        badge(s, ic, x + 0.1, my + 0.36 + (0.86 - 0.38) / 2, 0.38, bg=TEAL if best else NAVY)
        text(s, x + 0.56, my + 0.38, mw - 0.62, 0.82, [P(R(t, 10.5, True, TEAL if best else NAVY), space_after=1),
                                                       P(R(sub, 8.8, False, TEXT))],
             anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
        if i < 2:
            arrow(s, x + mw + 0.05, my + 0.36 + 0.33, 0.2, 0.2)
    # Validation protocol strip
    vy, vh = 5.38, 1.42
    card(s, RX, vy, RW, vh, fill=SKY, line=SKY_LINE, radius=0.1)
    section_label(s, RX + 0.14, vy + 0.1, RW - 0.28, "VALIDATION — against the real ocean, not just the training data",
                  ic="shield-check", size=10.5)
    facts = [("waves", "15,018", "Argo profiles, QC flags 1/2"), ("calendar-check", "2,639", "in the held-out 2023 test"),
             ("layers", "per depth", "RMSE · bias · r · skill vs climatology"), ("badge-check", "EN4", "independent Met Office cross-check")]
    fw = (RW - 0.28 - 3 * 0.12) / 4
    for i, (ic, big, small) in enumerate(facts):
        x = RX + 0.14 + i * (fw + 0.12)
        card(s, x, vy + 0.5, fw, 0.8, fill=WHITE, line=LINE, radius=0.07)
        icon(s, ic, x + 0.1, vy + 0.58, 0.24, TEAL)
        text(s, x + 0.4, vy + 0.53, fw - 0.44, 0.32, [P(R(big, 15, True, TEAL))], anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
        text(s, x + 0.1, vy + 0.87, fw - 0.16, 0.4, [P(R(small, 8.8, False, TEXT))], margin=0.01)


def slide4(s, rebuilt):
    clear_content_boxes(s, rebuilt)
    X, W = 0.35, 12.65
    section_label(s, X, 1.30, W, "Analysis of the feasibility of the idea",
                  "working prototype, validated on the held-out year 2023", ic="clipboard-check")
    cards = [
        ("waves", "2,639", "independent Argo profiles from 2023 — a year the model never trained on"),
        ("thermometer", "0.89 °C", "mean temperature error (RMSE, 15 depths) vs 1.10 °C for climatology"),
        ("arrow-down-to-line", "1.31 °C", "error at 100 m (thermocline) vs 1.80 °C for climatology"),
        ("crosshair", "70%", "of Argo values within ±1σ after calibration (ideal 68%; raw 41%)"),
    ]
    cg = 0.2
    cw = (W - 3 * cg) / 4
    for i, (ic, big, lab) in enumerate(cards):
        x = X + i * (cw + cg)
        card(s, x, 1.7, cw, 1.0, fill=NAVY, line=None, radius=0.1)
        badge(s, ic, x + 0.16, 1.8, 0.4, bg=TEAL)
        text(s, x + 0.66, 1.76, cw - 0.76, 0.48, [P(R(big, 22, True, WHITE))], anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
        text(s, x + 0.16, 2.26, cw - 0.3, 0.4, [P(R(lab, 9, False, ICE))], anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
    # Left column: feasible because
    top = 2.9
    LW = 3.85
    section_label(s, X, top, LW, "FEASIBLE BECAUSE", ic="circle-check", color=GREEN, bg=GREEN, size=11.5)
    feas = [
        ("database", "Data", "HIGH", "Every input, target and validation source used is open; Copernicus/ERA5 need only free sign-up"),
        ("cpu", "Technical", "HIGH", "Built and running end to end: ingestion → models → API → dashboard"),
        ("gauge", "Compute", "HIGH", "U-Net (1.09 M params) trained on a laptop CPU — no GPU"),
        ("indian-rupee", "Cost", "₹0", "No paid data or services used"),
    ]
    y, rh, pitch = top + 0.4, 0.62, 0.69
    for ic, name, rating, why in feas:
        card(s, X, y, LW, rh, fill=GREEN_LIGHT, line=GREEN_LINE, radius=0.08)
        icon(s, ic, X + 0.12, y + (rh - 0.28) / 2, 0.28, GREEN)
        text(s, X + 0.5, y + 0.06, 0.8, 0.24, [P(R(name, 10, True, NAVY))], anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
        box(s, X + 0.5, y + 0.34, 0.5, 0.2, fill=GREEN, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
        text(s, X + 0.5, y + 0.34, 0.5, 0.2, [P(R(rating, 8.5, True, WHITE), align=PP_ALIGN.CENTER)],
             anchor=MSO_ANCHOR.MIDDLE, margin=0)
        text(s, X + 1.36, y, LW - 1.44, rh, [P(R(why, 9, False, TEXT))], anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
        y += pitch
    sy = y + 0.04
    for tag, col, body in (("IMPLEMENTED", GREEN, "pipeline · models · validation · API · dashboard · Docker"),
                           ("PLANNED", AMBER, "hosted deployment · re-run on Copernicus/GLORYS with credentials")):
        box(s, X, sy + 0.02, 1.02, 0.2, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
        text(s, X, sy + 0.02, 1.02, 0.2, [P(R(tag, 7.5, True, WHITE), align=PP_ALIGN.CENTER)], anchor=MSO_ANCHOR.MIDDLE, margin=0)
        text(s, X + 1.1, sy - 0.02, LW - 1.1, 0.4, [P(R(body, 8.5))], margin=0.01)
        sy += 0.42
    # Right: challenges -> strategies
    CX = X + LW + 0.3
    CW = 3.6
    SX = CX + CW + 0.44
    SW = X + W - SX
    section_label(s, CX, top, CW, "Potential challenges and risks", ic="triangle-alert", color=AMBER, bg=AMBER, size=11.5)
    section_label(s, SX, top, SW, "Strategies for overcoming these challenges", ic="shield-check", color=TEAL, bg=TEAL, size=11.5)
    pairs = [
        ("Training target (a reanalysis) assimilates Argo → validation not fully independent",
         "Hold out a whole year; show the target product's own error as the ceiling (0.79 °C); cross-check against EN4"),
        ("Raw model uncertainty was over-confident (41% of values within ±1σ)",
         "Per-depth calibration fitted on 2022 floats, verified on 2023 → 70% within ±1σ, 94% within ±2σ"),
        ("Satellite salinity is gappy; removing it showed no measurable change in skill",
         "Merged SMAP + bias-corrected SMOS (absent days 209 → 1); re-test with Copernicus SSS"),
        ("GLORYS / Copernicus data need an account",
         "Credentialed adapters already built — switching source is a config change, no code change"),
        ("Thermocline (20–150 m) is the hardest layer to reconstruct",
         "Report error per depth, never one number; the uncertainty band widens where skill drops"),
    ]
    y = top + 0.4
    rh, pitch = 0.62, 0.7
    for ch, st in pairs:
        card(s, CX, y, CW, rh, fill=AMBER_LIGHT, line=AMBER_LINE, radius=0.08)
        icon(s, "triangle-alert", CX + 0.1, y + (rh - 0.24) / 2, 0.24, AMBER)
        text(s, CX + 0.42, y, CW - 0.48, rh, [P(R(ch, 9, False, TEXT))], anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
        arrow(s, CX + CW + 0.08, y + rh / 2 - 0.1, 0.28, 0.2)
        card(s, SX, y, SW, rh, fill=TEAL_LIGHT, line=TEAL_LINE, radius=0.08)
        icon(s, "shield-check", SX + 0.1, y + (rh - 0.24) / 2, 0.24, TEAL)
        text(s, SX + 0.42, y, SW - 0.48, rh, [P(R(st, 9, False, TEXT))], anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
        y += pitch


def slide5(s, shots, rebuilt):
    clear_content_boxes(s, rebuilt)
    X, LW = 0.35, 6.75
    section_label(s, X, 1.30, LW, "Potential impact on the target audience", ic="users")
    users = [("waves", "Ocean analysts &\nforecasters", "INCOIS-style"), ("tornado", "Cyclone\nforecasters", "heat available to storms"),
             ("fish", "Fisheries-advisory\nanalysts", "mixed layer & thermocline"), ("microscope", "Ocean\nresearchers", "daily 3-D record")]
    ug = 0.15
    uw = (LW - 3 * ug) / 4
    for i, (ic, u, sub) in enumerate(users):
        x = X + i * (uw + ug)
        card(s, x, 1.7, uw, 1.04, fill=NAVY, line=None, radius=0.1)
        icon(s, ic, x + (uw - 0.28) / 2, 1.78, 0.28, AQUA)
        text(s, x, 2.08, uw, 0.62, lines(u, 10, WHITE, True) + [P(R(sub, 8.5, False, ICE), align=PP_ALIGN.CENTER)],
             anchor=MSO_ANCHOR.MIDDLE, margin=0.02)
    text(s, X, 2.8, LW, 0.26, [P(R("Decision support — not a replacement for operational ocean models", 9.5, True, TEAL, italic=True))],
         margin=0.01)
    section_label(s, X, 3.14, LW, "Benefits of the solution", "social · economic · environmental · operational", ic="heart-handshake")
    tiles = [
        ("users", "SOCIAL", "Basin-wide cyclone heat potential (TCHP) and mixed-layer maps for every day and every 0.25° cell of the domain."),
        ("indian-rupee", "ECONOMIC", "Built only on free satellite data; the model trains on a CPU; products are precomputed, so viewing needs no heavy compute."),
        ("leaf", "ENVIRONMENTAL / CLIMATE", "A continuous daily record of upper-ocean heat and stratification for the North Indian Ocean (2019–2023)."),
        ("activity", "OPERATIONAL", "Each profile ships with a calibrated error bar and the nearest independent Argo float — trust is visible, not assumed."),
    ]
    tg = 0.15
    tw, th = (LW - tg) / 2, 1.16
    for i, (ic, t, body) in enumerate(tiles):
        x = X + (i % 2) * (tw + tg)
        y = 3.54 + (i // 2) * (th + 0.12)
        card(s, x, y, tw, th, fill=WHITE, line=SKY_LINE, radius=0.1, line_w=1.25)
        badge(s, ic, x + 0.14, y + 0.12, 0.38, bg=TEAL_LIGHT, fg=TEAL)
        text(s, x + 0.62, y + 0.12, tw - 0.74, 0.38, [P(R(t, 10.5, True, NAVY))], anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
        text(s, x + 0.14, y + 0.56, tw - 0.28, th - 0.6, [P(R(body, 9.3))], margin=0.01)
    # Key outcome bar
    box(s, X, 6.1, LW, 0.7, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    badge(s, "target", X + 0.16, 6.1 + 0.15, 0.4, bg=TEAL)
    text(s, X + 0.68, 6.1, LW - 0.8, 0.7, [P(R("KEY OUTCOME", 9, True, AQUA)),
                                          P(R("full-basin subsurface view · daily · 0–1000 m · with measured confidence", 11.5, True, WHITE))],
         anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
    # Right: prototype screenshot (cyclone fuel gauge)
    PX, PW = 7.4, 5.6
    section_label(s, PX, 1.30, PW, "Working prototype", "real event, held-out year", ic="monitor")
    half = (PW - 0.12) / 2
    _, _, ha = picture(s, shots / "fuel.png", PX, 1.7, w=half, crop=(0.40, 0.30, 0.67, 0.79))
    _, _, hb = picture(s, shots / "fuel.png", PX + half + 0.12, 1.7, w=half, crop=(0.68, 0.065, 1.0, 0.67))
    ph = max(ha, hb)
    capy = 1.7 + ph + 0.12
    card(s, PX, capy, PW, 6.1 - 0.14 - capy, fill=PANEL, line=LINE, radius=0.1)
    text(s, PX + 0.14, capy, PW - 0.28, 6.1 - 0.14 - capy, [P(
        R("Cyclone Fuel Gauge — ", 9.5, True, NAVY),
        R("reconstructed TCHP along Cyclone Mocha's real IBTrACS track (May 2023, ocean state 2 days before passage): "
          "52.7–102 kJ/cm² over the ocean, above the 50 kJ/cm² level commonly linked to intensification. "
          "Prototype output, not an operational forecast.", 9.5, False, MUTED))], anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
    chips = [("mouse-pointer-click", "Click any cell", "full 0–1000 m profile + σ"), ("map-pin", "Argo overlay", "nearest independent float"),
             ("file-down", "Export", "PDF / CSV report per point")]
    cg = 0.12
    cw = (PW - 2 * cg) / 3
    for i, (ic, t, sub) in enumerate(chips):
        x = PX + i * (cw + cg)
        card(s, x, 6.1, cw, 0.7, fill=TEAL_LIGHT, line=TEAL, radius=0.1, line_w=1.25)
        tw_ = text_w(t, 10, True)
        gx = x + (cw - (0.22 + 0.06 + tw_)) / 2
        icon(s, ic, gx, 6.1 + 0.12, 0.22, TEAL)
        text(s, gx + 0.28, 6.1 + 0.1, tw_ + 0.2, 0.26, [P(R(t, 10, True, TEAL))], anchor=MSO_ANCHOR.MIDDLE, margin=0)
        text(s, x, 6.1 + 0.38, cw, 0.24, [P(R(sub, 8.8, False, TEXT), align=PP_ALIGN.CENTER)], anchor=MSO_ANCHOR.MIDDLE, margin=0.01)


def slide6(s, rebuilt):
    clear_content_boxes(s, rebuilt)
    X, W = 0.35, 12.65
    section_label(s, X, 1.30, W, "Details / Links of the reference and research work", ic="book-open")
    refs = [
        ("SIH26066 OceanEmbed problem statement", "SIH · MoES / INCOIS", "Scope, domain, 15 standard depths", "sih.gov.in"),
        ("INCOIS Tropical Cyclone Heat Potential service", "INCOIS", "Operational relevance of TCHP", "incois.gov.in/site/services/tchp.jsp"),
        ("GLORYS12V1 global ocean reanalysis (PS-named target)", "Copernicus Marine", "Training-target specification", "doi.org/10.48670/moi-00021"),
        ("Argo float profiles via argopy", "Argo programme · Euro-Argo", "Independent validation", "argopy.readthedocs.io"),
        ("NOAA OISST v2.1 · NCEI Blended Seawinds", "NOAA NCEI", "Satellite inputs used", "ncei.noaa.gov"),
        ("EN4 quality-controlled ocean analysis", "UK Met Office Hadley Centre", "Independent cross-check", "metoffice.gov.uk/hadobs/en4"),
        ("Bay of Bengal barrier layer (Vinayachandran et al., 2002)", "JGR Oceans · AGU", "Why salinity is an input", "doi.org/10.1029/2001JC000831"),
        ("Subsurface temperature reconstruction, ConvLSTM vs LightGBM (2022)", "MDPI Remote Sensing", "Model choice; thermocline is hardest", "doi.org/10.3390/rs14133198"),
        ("IBTrACS v04r01 best-track archive", "NOAA NCEI", "Real cyclone tracks", "ncei.noaa.gov/products/international-best-track-archive"),
    ]
    rows, cols = len(refs) + 1, 5
    tbl = s.shapes.add_table(rows, cols, Inches(X), Inches(1.72), Inches(W), Inches(0.42 * rows)).table
    widths = [0.45, 4.6, 2.55, 2.4, 2.65]
    for i, w in enumerate(widths):
        tbl.columns[i].width = Inches(w)
    head = ["#", "Reference", "Organisation", "What it supports", "Link"]

    def cell(r, c, t, size=10, bold=False, color=TEXT, fill=None, align=PP_ALIGN.LEFT, link=None):
        ce = tbl.cell(r, c)
        ce.text = ""
        tf = ce.text_frame
        tf.word_wrap = True
        ce.margin_left = ce.margin_right = Inches(0.06)
        ce.margin_top = ce.margin_bottom = Inches(0.03)
        ce.vertical_anchor = MSO_ANCHOR.MIDDLE
        para = tf.paragraphs[0]
        para.alignment = align
        parts = t.replace("/products/", "/products/\n").split("\n") if link else [t]
        for k, part in enumerate(parts):
            if k:
                para.add_line_break()
            run = para.add_run()
            run.text = part
            run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color; set_font(run)
            if link:
                run.hyperlink.address = link
                run.font.color.rgb = color
        if fill is not None:
            ce.fill.solid(); ce.fill.fore_color.rgb = fill
    for c, h in enumerate(head):
        cell(0, c, h, 10.5, True, WHITE, NAVY, PP_ALIGN.CENTER if c == 0 else PP_ALIGN.LEFT)
    for r, ref in enumerate(refs, start=1):
        fill = WHITE if r % 2 else PANEL
        cell(r, 0, str(r), 10, True, TEAL, fill, PP_ALIGN.CENTER)
        cell(r, 1, ref[0], 10, True, NAVY, fill)
        cell(r, 2, ref[1], 9.5, False, TEXT, fill)
        cell(r, 3, ref[2], 9.5, False, TEXT, fill)
        cell(r, 4, ref[3], 9, False, BLUE, fill, link="https://" + ref[3])
    for r in range(rows):
        tbl.rows[r].height = Inches(0.42)
    # prototype link strip
    strip = card(s, X, 6.14, W, 0.64, fill=TEAL_LIGHT, line=TEAL_LINE, radius=0.1)
    strip.click_action.hyperlink.address = "https://github.com/saad-46/oceanembed"
    badge(s, "code-xml", X + 0.14, 6.14 + 0.13, 0.38, bg=NAVY)
    text(s, X + 0.64, 6.14, W - 0.8, 0.64, [P(
        R("Prototype source code & full research artifact:  ", 11, True, NAVY),
        R("github.com/saad-46/oceanembed", 12, True, BLUE),
        R("   ·   all reported metrics are generated by the pipeline (docs/RESULTS.md)", 10, False, MUTED))],
        anchor=MSO_ANCHOR.MIDDLE, margin=0.01)


def main():
    global ICON_DIR, LUCIDE_DIR
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True, help="official SIH template, or a previously built 6-slide deck")
    ap.add_argument("--shots", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--lucide", help="lucide-static/icons directory (only needed to render missing icons)")
    a = ap.parse_args()
    shots = Path(a.shots)
    ICON_DIR = shots / "icons"
    LUCIDE_DIR = Path(a.lucide) if a.lucide else None
    prs = Presentation(a.template)
    s = list(prs.slides)
    rebuilt = len(s) == 6  # already-built deck: instructions slide gone, content shapes to rebuild
    slide1(s[0])
    slide2(s[1], shots, rebuilt)
    slide3(s[2], rebuilt)
    slide4(s[3], rebuilt)
    slide5(s[4], shots, rebuilt)
    slide6(s[5], rebuilt)
    for sl in s[1:6]:
        set_team_oval(sl)
    if not rebuilt:
        delete_slide(prs, 6)  # "IMPORTANT INSTRUCTIONS" — must not be uploaded
    prs.save(a.out)
    print("saved", a.out, "slides:", len(prs.slides))


if __name__ == "__main__":
    main()
